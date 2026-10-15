"""Regenerate the resume test fixtures (fictional people, no real data):

    cd backend
    .venv\\Scripts\\python.exe tests\\fixtures\\make_fixtures.py

Needs a TrueType font (Arial on Windows) only for the two image fixtures. The generated files are committed,
so the tests themselves need no fonts.
"""
from pathlib import Path

HERE = Path(__file__).resolve().parent

BACKEND_RESUME = """AARAV MEHTA
aarav.mehta@example.com | +91 98765 43210
linkedin.com/in/aarav-mehta | github.com/aarav-mehta

SUMMARY
Final year B.Tech student interested in backend engineering.

SKILLS
Languages: Python, Java, C, SQL
Frameworks: FastAPI, Django
Databases: MongoDB, PostgreSQL
Tools: Git, Docker, Linux

PROJECTS
Library Management API | FastAPI, MongoDB
- Built a REST API with FastAPI and MongoDB serving 500 requests per minute
- Implemented JWT authentication with 3 user roles
- Deployed the service with Docker on AWS EC2
Placement Portal | Django, PostgreSQL
- Designed 12 database tables and wrote 25 SQL queries for reports
- Worked on the login page and some other features

INTERNSHIP
Backend Intern, Acme Software, Jun 2025 - Aug 2025
- Reduced API response time by 35% by adding Redis caching
- Wrote unit tests with pytest, raising coverage to 80%

EDUCATION
B.Tech in Computer Science, Example Institute of Technology, 2022-2026, CGPA 8.4/10

CERTIFICATIONS
AWS Cloud Practitioner (2025)
"""

CLOUD_RESUME = """MEERA IYER
meera.iyer@example.com | +91 91234 56780
linkedin.com/in/meera-iyer | github.com/meera-iyer

SKILLS
Cloud: AWS, GCP
DevOps: Docker, Kubernetes, Terraform, Linux, CI/CD, Prometheus, Grafana
Languages: Python, Bash

PROJECTS
Kubernetes Deployment Pipeline
- Automated build and deployment of 4 microservices using GitHub Actions and Docker
- Provisioned an EKS cluster with Terraform, cutting setup time from 2 days to 40 minutes
- Configured Prometheus and Grafana dashboards to monitor 12 services

EXPERIENCE
Cloud Intern, Nimbus Labs, Jan 2025 - Jun 2025
- Migrated 6 virtual machines to AWS and documented the Linux hardening steps
- Reduced monthly cloud cost by 18% by right-sizing instances

EDUCATION
B.E. in Information Technology, Example College of Engineering, 2021-2025, CGPA 8.1/10

CERTIFICATIONS
AWS Solutions Architect Associate (2025)
Computer Networks and TCP/IP fundamentals
"""


def render_image(text: str, width: int = 1240, height: int = 1754):
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 27)
    big = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 38)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    y = 60
    for i, line in enumerate(text.split("\n")):
        draw.text((70, y), line, fill=(20, 20, 20), font=big if i == 0 else font)
        y += 52 if i == 0 else 38
    return img


def make_text_pdf(text: str, path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 60
    for i, line in enumerate(text.split("\n")):
        page.insert_text((50, y), line, fontsize=16 if i == 0 else 10, fontname="helv")
        y += 22 if i == 0 else 14
    doc.save(path)


def make_scanned_pdf(text: str, path: Path) -> None:
    import io

    import fitz

    buf = io.BytesIO()
    render_image(text).save(buf, "PNG")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=buf.getvalue())  # an image and no text layer
    doc.save(path)


def make_docx(text: str, path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Aarav Mehta - Resume"
    lines = text.split("\n")
    skills_at = lines.index("SKILLS")
    for line in lines[:skills_at + 1]:
        doc.add_paragraph(line)
    table = doc.add_table(rows=0, cols=2)
    for line in lines[skills_at + 1:skills_at + 5]:
        label, rest = line.split(": ", 1)
        row = table.add_row().cells
        row[0].text, row[1].text = label + ":", rest
    for line in lines[skills_at + 5:]:
        doc.add_paragraph(line)
    doc.save(path)


if __name__ == "__main__":
    make_text_pdf(BACKEND_RESUME, HERE / "resume_text.pdf")
    make_scanned_pdf(BACKEND_RESUME, HERE / "resume_scanned.pdf")
    render_image(BACKEND_RESUME).save(HERE / "resume_photo.jpg", "JPEG", quality=82)
    make_docx(BACKEND_RESUME, HERE / "resume.docx")
    (HERE / "resume.txt").write_bytes(BACKEND_RESUME.encode("utf-8"))
    make_text_pdf(CLOUD_RESUME, HERE / "resume_cloud.pdf")
    print("fixtures written:", ", ".join(sorted(p.name for p in HERE.iterdir() if p.name != "make_fixtures.py")))
